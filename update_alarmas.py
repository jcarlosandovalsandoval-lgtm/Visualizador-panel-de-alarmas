#!/usr/bin/env python3
"""
Actualiza data/alarmas.json a partir de los archivos Excel (.xlsx) y de
texto (.txt) en uploads/.

Uso:
    python3 update_alarmas.py

Cada vez que se ejecuta:
  - Lee todos los .xlsx y .txt dentro de uploads/
  - Extrae fecha, hora, dia y evento:
      * En .xlsx busca la hoja/columnas por nombre, asi que acepta
        encabezados en cualquier orden o con nombres levemente distintos
        ("Fecha", "Hora", "Dia"/"Día", "Evento").
      * En .txt reconoce el formato de impresion del panel
        ("ENTRY <n>  HH:MM:SS  DIA  DD-MON-AA  <evento>", con lineas de
        continuacion para eventos largos) o texto tabulado/delimitado con
        columnas fecha/hora/dia/evento.
  - Clasifica cada evento (tipo) y su ubicacion
  - Combina con lo que ya hay en data/alarmas.json, sin duplicar filas
    (una fila se considera duplicada si coincide fecha+hora+evento)
  - Reordena cronologicamente y reescribe data/alarmas.json

Es seguro correrlo varias veces sobre los mismos archivos: no genera
duplicados.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
UPLOADS_DIR = ROOT / "uploads"
DATA_PATH = ROOT / "data" / "alarmas.json"

LOC_MARKERS = ["SMOKE DETECTOR", "PULL STATION", "FIRE MONITOR ZONE", "FIRE ALARM"]
TIPO_PATTERNS = [
    ("ALARM SILENCE", "silenciada"),
    ("TEST ABNORMAL", "test"),
    ("DATA NOT AVAILABLE", "falla"),
    ("ACKNOWLEDGED", "reconocida"),
    ("ACK REQUESTS", "reconocida"),
    ("FIRE ALARM", "alarma"),
    ("RESET REQUESTED", "reset"),
]

HEADER_ALIASES = {
    "fecha": "fecha",
    "hora": "hora",
    "dia": "dia",
    "día": "dia",
    "evento": "evento",
}


def classify_tipo(evento):
    for pat, tipo in TIPO_PATTERNS:
        if pat in evento:
            return tipo
    return "otro"


def extract_ubicacion(evento, tipo):
    if tipo not in ("falla", "alarma", "test"):
        return None
    best = None
    for marker in LOC_MARKERS:
        idx = evento.find(marker)
        if idx > 0 and (best is None or idx < best):
            best = idx
    if best is None:
        return None
    loc = evento[:best].strip()
    return loc or None


def find_header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=10):
        values = [str(c.value).strip().lower() if c.value is not None else "" for c in row]
        if any(v in HEADER_ALIASES for v in values):
            col_map = {}
            for idx, v in enumerate(values):
                if v in HEADER_ALIASES:
                    col_map[HEADER_ALIASES[v]] = idx
            if "fecha" in col_map and "evento" in col_map:
                return row[0].row, col_map
    return None, None


def parse_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows_out = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row, col_map = find_header_row(ws)
        if not col_map:
            continue
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            fecha_val = row[col_map["fecha"]] if col_map.get("fecha") is not None else None
            evento_val = row[col_map["evento"]] if col_map.get("evento") is not None else None
            if fecha_val is None or evento_val is None:
                continue
            evento = str(evento_val).strip()
            if not evento:
                continue

            fecha = fecha_val.strftime("%Y-%m-%d") if hasattr(fecha_val, "strftime") else str(fecha_val)

            hora_val = row[col_map["hora"]] if col_map.get("hora") is not None else None
            hora = hora_val.strftime("%H:%M:%S") if hasattr(hora_val, "strftime") else (str(hora_val) if hora_val else "")

            dia_val = row[col_map["dia"]] if col_map.get("dia") is not None else None
            dia = str(dia_val).strip() if dia_val else ""

            tipo = classify_tipo(evento)
            ubicacion = extract_ubicacion(evento, tipo)

            rows_out.append({
                "fecha": fecha,
                "hora": hora,
                "dia": dia,
                "evento": evento,
                "tipo": tipo,
                "ubicacion": ubicacion,
            })
    return rows_out


MESES = {
    "ENE": 1, "JAN": 1, "FEB": 2, "MAR": 3, "ABR": 4, "APR": 4, "MAY": 5,
    "JUN": 6, "JUL": 7, "AGO": 8, "AUG": 8, "SEP": 9, "SET": 9,
    "OCT": 10, "NOV": 11, "DIC": 12, "DEC": 12,
}

ENTRY_RE = re.compile(
    r"ENTRY\s+(?P<num>\d+)\s+(?P<hora>\d{1,2}:\d{2}:\d{2})\s+(?P<dia>[A-Z]{3})\s+"
    r"(?P<dd>\d{1,2})-(?P<mon>[A-Z]{3})-(?P<yy>\d{2,4})\s*(?P<evento>.*)",
    re.IGNORECASE,
)

DATE_ISO_RE = re.compile(r"^(?P<y>\d{4})-(?P<m>\d{1,2})-(?P<d>\d{1,2})$")
DATE_DMY_DASH_RE = re.compile(r"^(?P<d>\d{1,2})-(?P<mon>[A-Za-z]{3})-(?P<y>\d{2,4})$")
DATE_DMY_SLASH_RE = re.compile(r"^(?P<d>\d{1,2})/(?P<m>\d{1,2})/(?P<y>\d{2,4})$")


def normalize_year(yy):
    yy = int(yy)
    return yy if yy > 100 else 2000 + yy


def parse_fecha_str(s):
    s = s.strip()
    m = DATE_ISO_RE.match(s)
    if m:
        return f"{int(m['y']):04d}-{int(m['m']):02d}-{int(m['d']):02d}"
    m = DATE_DMY_DASH_RE.match(s)
    if m:
        mon = MESES.get(m["mon"].upper())
        if mon:
            return f"{normalize_year(m['y']):04d}-{mon:02d}-{int(m['d']):02d}"
    m = DATE_DMY_SLASH_RE.match(s)
    if m:
        return f"{normalize_year(m['y']):04d}-{int(m['m']):02d}-{int(m['d']):02d}"
    return None


def parse_hora_str(s):
    s = s.strip()
    parts = s.split(":")
    if len(parts) == 2:
        s = s + ":00"
    if re.match(r"^\d{1,2}:\d{2}:\d{2}$", s):
        h, mi, se = s.split(":")
        return f"{int(h):02d}:{int(mi):02d}:{int(se):02d}"
    return None


BOILERPLATE_RE = re.compile(
    r"^-+$|^service port\b|^report\s+\d+\s*:|press return for next screen|ctrl-x to abort|"
    r"^alarm historical log report completed",
    re.IGNORECASE,
)


def is_boilerplate(line):
    return bool(BOILERPLATE_RE.search(line))


def parse_txt_entry_format(lines):
    entries = []
    current = None
    for raw_line in lines:
        line = raw_line.rstrip("\n").strip()
        if not line or is_boilerplate(line):
            continue
        m = ENTRY_RE.match(line)
        if m:
            if current:
                entries.append(current)
            mon = MESES.get(m["mon"].upper())
            fecha = f"{normalize_year(m['yy']):04d}-{mon:02d}-{int(m['dd']):02d}" if mon else None
            hora = parse_hora_str(m["hora"])
            current = {
                "fecha": fecha,
                "hora": hora,
                "dia": m["dia"].upper(),
                "evento": m["evento"].strip(),
            }
        elif current is not None:
            current["evento"] = (current["evento"] + " " + line).strip()
    if current:
        entries.append(current)
    for e in entries:
        e["evento"] = re.sub(r"\s+", " ", e["evento"]).strip()
    return [e for e in entries if e["fecha"] and e["hora"] and e["evento"]]


def split_delimited_line(line):
    if "\t" in line:
        return [p.strip() for p in line.split("\t")]
    if re.search(r" {2,}", line):
        return [p.strip() for p in re.split(r" {2,}", line) if p.strip()]
    if "," in line:
        return [p.strip() for p in line.split(",")]
    return None


def parse_txt_delimited_format(lines):
    rows = [split_delimited_line(l.rstrip("\n")) for l in lines if l.strip()]
    rows = [r for r in rows if r]
    if not rows:
        return []

    header = [c.strip().lower() for c in rows[0]]
    col_map = {}
    for idx, name in enumerate(header):
        if name in HEADER_ALIASES:
            col_map[HEADER_ALIASES[name]] = idx
    data_rows = rows[1:] if "fecha" in col_map and "evento" in col_map else rows
    if not col_map:
        n = len(rows[0])
        if n == 5:
            col_map = {"fecha": 1, "hora": 2, "dia": 3, "evento": 4}
        elif n == 4:
            col_map = {"fecha": 0, "hora": 1, "dia": 2, "evento": 3}
        else:
            return []

    entries = []
    for r in data_rows:
        try:
            fecha_raw = r[col_map["fecha"]]
            evento = r[col_map["evento"]]
        except IndexError:
            continue
        fecha = parse_fecha_str(fecha_raw)
        if not fecha or not evento.strip():
            continue
        hora = parse_hora_str(r[col_map["hora"]]) if "hora" in col_map and col_map["hora"] < len(r) else ""
        dia = r[col_map["dia"]].strip() if "dia" in col_map and col_map["dia"] < len(r) else ""
        entries.append({"fecha": fecha, "hora": hora or "", "dia": dia, "evento": evento.strip()})
    return entries


def parse_txt(path):
    text = path.read_text(encoding="utf-8", errors="ignore")
    lines = text.splitlines()

    raw_entries = parse_txt_entry_format(lines)
    if not raw_entries:
        raw_entries = parse_txt_delimited_format(lines)

    rows_out = []
    for e in raw_entries:
        tipo = classify_tipo(e["evento"])
        ubicacion = extract_ubicacion(e["evento"], tipo)
        rows_out.append({
            "fecha": e["fecha"],
            "hora": e["hora"],
            "dia": e["dia"],
            "evento": e["evento"],
            "tipo": tipo,
            "ubicacion": ubicacion,
        })
    return rows_out


def dedup_key(entry):
    return (entry["fecha"], entry["hora"], entry["evento"])


def main():
    if DATA_PATH.exists():
        existing = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    else:
        existing = []

    seen = {dedup_key(e) for e in existing}
    combined = list(existing)

    added = 0
    xlsx_files = sorted(UPLOADS_DIR.glob("*.xlsx"))
    txt_files = sorted(UPLOADS_DIR.glob("*.txt"))
    if not xlsx_files and not txt_files:
        print(f"No se encontraron archivos .xlsx ni .txt en {UPLOADS_DIR}")
        return

    for path, parser in [(p, parse_xlsx) for p in xlsx_files if not p.name.startswith("~$")] + \
                         [(p, parse_txt) for p in txt_files]:
        try:
            new_rows = parser(path)
        except Exception as exc:
            print(f"  ! No se pudo leer {path.name}: {exc}", file=sys.stderr)
            continue
        file_added = 0
        for entry in new_rows:
            key = dedup_key(entry)
            if key in seen:
                continue
            seen.add(key)
            combined.append(entry)
            file_added += 1
            added += 1
        print(f"  {path.name}: {len(new_rows)} filas leidas, {file_added} nuevas")

    combined.sort(key=lambda e: (e["fecha"], e["hora"]))
    for i, entry in enumerate(combined, start=1):
        entry["id"] = i
        # keep field order stable: id, fecha, hora, dia, evento, tipo, ubicacion
        ordered = {
            "id": entry["id"],
            "fecha": entry["fecha"],
            "hora": entry["hora"],
            "dia": entry["dia"],
            "evento": entry["evento"],
            "tipo": entry["tipo"],
            "ubicacion": entry["ubicacion"],
        }
        combined[i - 1] = ordered

    DATA_PATH.parent.mkdir(parents=True, exist_ok=True)
    DATA_PATH.write_text(json.dumps(combined, ensure_ascii=False, indent=2), encoding="utf-8")

    if added:
        print(f"Listo: {added} eventos nuevos agregados. Total en la base: {len(combined)}.")
    else:
        print(f"Sin cambios: no hay eventos nuevos. Total en la base: {len(combined)}.")


if __name__ == "__main__":
    main()
